"""從舊領據 Word 檔讀取個人資料，建立姓名→ReceiptInfo 的查找表

支援：
- .docx（python-docx 直讀）
- .doc（win32com 先轉成暫存 .docx，再用 python-docx 讀）

用法：
    lookup = load_receipts_from_dir("D:/115年3月-領據/")
    # lookup == {"王永良": ReceiptInfo(id_number="A123...", address="台北...", ...)}
"""

import os
import tempfile
from docx import Document
from docx.oxml.ns import qn

from models import ReceiptInfo

# 對應表：中文標籤關鍵字 → ReceiptInfo 欄位
LABEL_FIELD_MAP = [
    ("具領",    "recipient_name"),
    ("身分證",  "id_number"),
    ("戶籍",    "address"),
    ("聯絡電話", "phone"),
    ("戶名",    "account_name"),
    ("銀行及分行", "bank_branch"),
    ("銀行代碼", "bank_code"),
    ("帳號",    "account_number"),
]


def load_receipts_from_dir(dir_path: str) -> dict[str, ReceiptInfo]:
    """掃描資料夾，讀取所有 *核銷領據.doc/.docx，
    回傳 {姓名: ReceiptInfo}（金額欄位設為 0，由主程式填入）
    """
    if not dir_path or not os.path.isdir(dir_path):
        return {}

    lookup: dict[str, ReceiptInfo] = {}

    for fname in os.listdir(dir_path):
        base, ext = os.path.splitext(fname)
        ext_lower = ext.lower()
        if ext_lower not in (".doc", ".docx"):
            continue
        if "領據" not in base and "receipt" not in base.lower():
            continue

        # 從檔名推出姓名（取掉「核銷領據」等後綴）
        name = _extract_name_from_filename(base)
        if not name:
            continue

        fpath = os.path.join(dir_path, fname)
        try:
            info = _extract_receipt_info(fpath, ext_lower)
            if info:
                info.recipient_name = info.recipient_name or name
                lookup[name] = info
                print(f"  [領據] 讀入 {fname} → {name}")
        except Exception as e:
            print(f"  [WARN] 讀取領據失敗: {fname} - {e}")

    return lookup


def _extract_name_from_filename(base: str) -> str:
    """從檔名去掉常見後綴，取得人名
    例：王永良核銷領據 → 王永良
        洪德仁核銷領據(扣稅) → 洪德仁
    """
    # 去掉括號內容
    import re
    base = re.sub(r"[（(][^）)]*[）)]", "", base).strip()
    # 常見後綴
    for suffix in ("核銷領據", "核销领据", "領據", "领据"):
        if base.endswith(suffix):
            return base[: -len(suffix)].strip()
    # 找最後一個常見詞之前的部分
    for kw in ("核銷", "核销"):
        idx = base.rfind(kw)
        if idx > 0:
            return base[:idx].strip()
    # 如果都沒找到，就用整個 base（可能本身就是姓名）
    return base.strip()


def _extract_receipt_info(fpath: str, ext: str) -> ReceiptInfo | None:
    """讀取一份領據檔，抽取個人資料"""
    if ext == ".docx":
        doc = Document(fpath)
    elif ext == ".doc":
        doc = _doc_to_docx_and_open(fpath)
    else:
        return None

    if doc is None:
        return None

    info = ReceiptInfo()
    _fill_from_tables(doc, info)
    _fill_from_paragraphs(doc, info)
    return info


def _doc_to_docx_and_open(doc_path: str):
    """用 win32com 將 .doc 轉成暫存 .docx，再用 python-docx 開啟"""
    try:
        import win32com.client
    except ImportError:
        print("  [WARN] win32com 未安裝，無法讀取 .doc 檔")
        return None

    tmp_path = os.path.join(
        tempfile.gettempdir(),
        f"receipt_tmp_{os.getpid()}.docx"
    )
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wdoc = word.Documents.Open(os.path.abspath(doc_path))
        wdoc.SaveAs(os.path.abspath(tmp_path), FileFormat=12)  # 12 = docx
        wdoc.Close()
    except Exception as e:
        print(f"  [WARN] win32com 轉換失敗: {e}")
        return None
    finally:
        word.Quit()

    try:
        return Document(tmp_path)
    except Exception:
        return None
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


def _fill_from_tables(doc, info: ReceiptInfo):
    """從表格的每一行掃描 標籤|值 結構"""
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 2:
                continue
            label_text = _cell_text(cells[0])
            value_text = _cell_text(cells[1])
            _try_assign(info, label_text, value_text)


def _fill_from_paragraphs(doc, info: ReceiptInfo):
    """從 body 段落掃描「標籤：值」格式"""
    for p in doc.paragraphs:
        text = p.text.strip()
        for sep in ("：", ":"):
            if sep in text:
                parts = text.split(sep, 1)
                if len(parts) == 2:
                    _try_assign(info, parts[0].strip(), parts[1].strip())
                break


def _try_assign(info: ReceiptInfo, label: str, value: str):
    """根據 label 文字判斷要填哪個欄位"""
    value = value.strip()
    if not value:
        return
    for keyword, field in LABEL_FIELD_MAP:
        if keyword in label:
            # 只填入尚未設定的欄位（避免重複覆蓋）
            if not getattr(info, field):
                setattr(info, field, value)
            return


def _cell_text(cell) -> str:
    """取得 cell 的完整文字"""
    return "".join(
        t.text for t in cell._tc.iter(qn("w:t"))
    ).strip()
