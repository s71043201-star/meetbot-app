"""產生按醫師/診所分 sheet 的 Excel 統計檔

格式完全對照原始系統匯出的 Excel：
- 字型：標楷體 12pt，不粗體，黑色
- 對齊：全部置中 + 垂直置中
- 標題列 merged（每欄 row 1-3 合併）
- 行高：標題 15，資料 16.15
- 所有 cell 有 thin 框線
- 摘要行 merged，不粗體
"""

import os
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CELL_FONT = Font(name="標楷體", size=12, bold=False, color="FF000000")
CELL_ALIGN_WRAP = Alignment(horizontal="center", vertical="center",
                            wrap_text=True)
CELL_ALIGN = Alignment(horizontal="center", vertical="center")

HEADER_ROW_HEIGHT = 15.0
DATA_ROW_HEIGHT = 16.15

COL_NAME = 1
COL_BIRTH = 3
COL_PTYPE = 6
COL_CLINIC = 7
COL_DOCTOR = 8
COL_EXEC_DONE = 14
COL_EXEC_DATE = 16
COL_DATE = 21

FEE_PRESCRIPTION = 300
FEE_EXECUTION = 100
FEE_HEALTH_MGMT = 7000

# 處方類型排序順序
PTYPE_ORDER = {"營養處方": 0, "運動處方": 1, "情緒調適處方": 2, "社會處方": 3}


def _sort_by_ptype(rows):
    """按處方類型排序"""
    return sorted(rows, key=lambda r: PTYPE_ORDER.get(str(r[COL_PTYPE] or ""), 99))


def _setup_sheet(ws, headers, col_widths, wrap_cols=None):
    """設定 sheet：merged 標題列 + 欄寬 + 行高"""
    if wrap_cols is None:
        wrap_cols = {0, 1, 2}  # A, B, C default wrap

    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = CELL_FONT
        cell.border = CELL_BORDER
        cell.alignment = CELL_ALIGN_WRAP if (c - 1) in wrap_cols else CELL_ALIGN
        # Merge row 1-3
        col_letter = get_column_letter(c)
        ws.merge_cells(f"{col_letter}1:{col_letter}3")

    # Row heights for header + borders on ALL cells in merged header area
    for r in range(1, 4):
        ws.row_dimensions[r].height = HEADER_ROW_HEIGHT
    # openpyxl merged cell 框線需要逐 cell 設定才能正確顯示
    for c in range(1, len(headers) + 1):
        for r in range(1, 4):
            cell = ws.cell(r, c)
            cell.border = CELL_BORDER
            cell.font = CELL_FONT
            cell.alignment = CELL_ALIGN_WRAP

    # Column widths
    for c, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def _write_row(ws, row_num, values, wrap_cols=None):
    """寫入一列，所有 cell 置中 + 框線 + 標楷體 12pt"""
    if wrap_cols is None:
        wrap_cols = {0, 1}  # A, B default wrap

    ws.row_dimensions[row_num].height = DATA_ROW_HEIGHT
    for c, v in enumerate(values, 1):
        cell = ws.cell(row_num, c, v)
        cell.font = CELL_FONT
        cell.border = CELL_BORDER
        cell.alignment = CELL_ALIGN_WRAP if (c - 1) in wrap_cols else CELL_ALIGN

    # 補齊到 F 欄（若 values 不足 6 欄），確保所有欄都有框線
    for c in range(len(values) + 1, 7):
        cell = ws.cell(row_num, c)
        cell.border = CELL_BORDER


def _write_summary(ws, row_num, text, num_cols=6):
    """摘要行（merged，不粗體，有框線）"""
    cell = ws.cell(row_num, 1, text)
    cell.font = CELL_FONT
    cell.border = CELL_BORDER
    ws.merge_cells(
        start_row=row_num, start_column=1,
        end_row=row_num, end_column=num_cols
    )
    # 框線補到所有 cell
    for c in range(2, num_cols + 1):
        ws.cell(row_num, c).border = CELL_BORDER


def _write_signature(ws, row_num, num_cols=6, span_rows=5):
    """簽章行（merged 多行，有框線）"""
    cell = ws.cell(row_num, 1, "確認簽章:")
    cell.font = CELL_FONT
    cell.border = CELL_BORDER
    ws.merge_cells(
        start_row=row_num, start_column=1,
        end_row=row_num + span_rows - 1, end_column=num_cols
    )
    # 框線補到所有 merged cell
    for r in range(row_num, row_num + span_rows):
        for c in range(1, num_cols + 1):
            ws.cell(r, c).border = CELL_BORDER


def generate_prescription_fee_excel(records: list, prefix: str,
                                    output_dir: str):
    """處方開立費 Excel"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    by_doctor = defaultdict(list)
    for row in records:
        if row[COL_PTYPE] and row[COL_DOCTOR]:
            by_doctor[str(row[COL_DOCTOR])].append(row)

    headers = ["序號", "民眾姓名", "民眾出生年月日(yyyy/mm/dd)",
               "處方類型", "開立醫師", "開立日期時間"]
    col_widths = [8.57, 13.86, 18.14, 14.14, 12.71, 26.0]

    for doctor, rows in sorted(by_doctor.items()):
        ws = wb.create_sheet(f"處方開立-{doctor}"[:31])
        _setup_sheet(ws, headers, col_widths, wrap_cols={0, 1, 2})

        sorted_rows = _sort_by_ptype(rows)
        for seq, row in enumerate(sorted_rows, 1):
            _write_row(ws, seq + 3, [
                seq,
                row[COL_NAME],
                str(row[COL_BIRTH] or ""),
                str(row[COL_PTYPE] or ""),
                doctor,
                str(row[COL_DATE] or ""),
            ], wrap_cols={0, 1})

        n = len(sorted_rows)
        summary_row = n + 4
        _write_summary(ws, summary_row,
                       f"{prefix}處方開立份數：{n}份")
        _write_summary(ws, summary_row + 1,
                       f"總申報金額（元）：{n * FEE_PRESCRIPTION:,}")

    path = os.path.join(output_dir, f"處方費用-處方開立費-{prefix}-總.xlsx")
    wb.save(path)
    return path


def generate_execution_fee_excel(records: list, prefix: str,
                                 output_dir: str):
    """處方執行費 Excel"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    by_doctor = defaultdict(list)
    for row in records:
        if row[COL_EXEC_DONE] and row[COL_DOCTOR]:
            by_doctor[str(row[COL_DOCTOR])].append(row)

    headers = ["序號", "民眾姓名", "民眾出生年月日(yyyy/mm/dd)",
               "處方類型", "處方開立醫師", "執行日期"]
    col_widths = [8.57, 13.86, 18.14, 14.14, 14.14, 26.0]

    for doctor, rows in sorted(by_doctor.items()):
        ws = wb.create_sheet(f"處方執行-{doctor}"[:31])
        _setup_sheet(ws, headers, col_widths, wrap_cols={0, 1, 2})

        sorted_rows = _sort_by_ptype(rows)
        for seq, row in enumerate(sorted_rows, 1):
            _write_row(ws, seq + 3, [
                seq,
                row[COL_NAME],
                str(row[COL_BIRTH] or ""),
                str(row[COL_PTYPE] or ""),
                doctor,
                str(row[COL_EXEC_DATE] or ""),
            ], wrap_cols={0, 1})

        n = len(sorted_rows)
        summary_row = n + 4
        _write_summary(ws, summary_row,
                       f"{prefix}完成執行處方份數：{n}份")
        _write_summary(ws, summary_row + 1,
                       f"處方執行費總申報金額（元）：{n * FEE_EXECUTION:,}")

    path = os.path.join(output_dir, f"處方費用-處方執行費-{prefix}-總.xlsx")
    wb.save(path)
    return path


def generate_health_mgmt_excel(records: list, prefix: str,
                               output_dir: str):
    """健康管理費 Excel"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    by_clinic = defaultdict(list)
    for row in records:
        if row[COL_CLINIC]:
            by_clinic[str(row[COL_CLINIC])].append(row)

    headers = ["序號", "民眾姓名", "民眾出生年月日(yyyy/mm/dd)",
               "處方類型", "協助處方開立行政人員", "開立日期時間"]
    col_widths = [6.86, 12.57, 17.43, 15.86, 17.43, 26.0]

    for clinic, rows in sorted(by_clinic.items()):
        ws = wb.create_sheet(f"健康管理費-{clinic}"[:31])
        _setup_sheet(ws, headers, col_widths, wrap_cols={0, 1, 2})

        sorted_rows = _sort_by_ptype(rows)
        for seq, row in enumerate(sorted_rows, 1):
            _write_row(ws, seq + 3, [
                seq,
                row[COL_NAME],
                str(row[COL_BIRTH] or ""),
                str(row[COL_PTYPE] or ""),
                str(row[COL_DOCTOR] or ""),
                str(row[COL_DATE] or ""),
            ], wrap_cols={0, 1})

        n = len(sorted_rows)
        summary_row = n + 4
        _write_summary(ws, summary_row,
                       f"健康管理費總申報金額（元）：{FEE_HEALTH_MGMT:,}")

    path = os.path.join(output_dir, f"健康管理費-{prefix}-總.xlsx")
    wb.save(path)
    return path


def generate_health_mgmt_excel_per_clinic(records: list, prefix: str,
                                           output_dir: str,
                                           clinic_to_person: dict | None = None):
    """健康管理費 Excel — 每間診所各自一份 xlsx（檔名用 clinic_person 或 clinic 名）

    clinic_to_person: {診所名: 診所人員姓名}，若提供則用人員名作檔名
    return: list of created xlsx paths
    """
    headers = ["序號", "民眾姓名", "民眾出生年月日(yyyy/mm/dd)",
               "處方類型", "協助處方開立行政人員", "開立日期時間"]
    col_widths = [6.86, 12.57, 17.43, 15.86, 17.43, 26.0]

    by_clinic = defaultdict(list)
    for row in records:
        if row[COL_CLINIC]:
            by_clinic[str(row[COL_CLINIC])].append(row)

    paths = []
    os.makedirs(output_dir, exist_ok=True)
    for clinic, rows in sorted(by_clinic.items()):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"健康管理費-{clinic}"[:31]
        _setup_sheet(ws, headers, col_widths, wrap_cols={0, 1, 2})

        sorted_rows = _sort_by_ptype(rows)
        for seq, row in enumerate(sorted_rows, 1):
            _write_row(ws, seq + 3, [
                seq,
                row[COL_NAME],
                str(row[COL_BIRTH] or ""),
                str(row[COL_PTYPE] or ""),
                str(row[COL_DOCTOR] or ""),
                str(row[COL_DATE] or ""),
            ], wrap_cols={0, 1})

        n = len(sorted_rows)
        summary_row = n + 4
        _write_summary(ws, summary_row,
                       f"健康管理費總申報金額（元）：{FEE_HEALTH_MGMT:,}")

        # 檔名：優先用 person 名，否則 clinic
        person = (clinic_to_person or {}).get(clinic, "") or clinic
        # 清掉不能做檔名的字元
        person_safe = person.replace("/", "-").replace("\\", "-")
        path = os.path.join(output_dir, f"{person_safe}_健康管理費-{prefix}.xlsx")
        wb.save(path)
        paths.append(path)

    return paths


def _xlsx_to_pdf(xlsx_path: str) -> str | None:
    """用 Excel COM 將 xlsx 轉為 PDF。失敗 return None
    會自動設橫向 + 縮放到 1 頁寬。
    """
    try:
        import win32com.client
    except ImportError:
        return None
    pdf_path = xlsx_path.rsplit(".", 1)[0] + ".pdf"
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        try:
            excel.Visible = False
        except Exception:
            pass
        try:
            excel.DisplayAlerts = False
        except Exception:
            pass
        try:
            wb = excel.Workbooks.Open(os.path.abspath(xlsx_path))
            _apply_pdf_pagesetup(wb)
            # 0 = xlTypePDF
            wb.ExportAsFixedFormat(0, os.path.abspath(pdf_path))
            wb.Close(SaveChanges=False)
            return pdf_path
        finally:
            try:
                excel.Quit()
            except Exception:
                pass
    except Exception:
        return None


def _apply_pdf_pagesetup(wb):
    """設定每個 sheet 的 PageSetup：橫向 + 縮放到 1 頁寬（高度不限），小邊距"""
    for sheet_idx in range(1, wb.Sheets.Count + 1):
        ws = wb.Sheets(sheet_idx)
        try:
            ps = ws.PageSetup
            ps.Orientation = 2        # xlLandscape
            ps.PaperSize = 9          # xlPaperA4
            ps.Zoom = False           # 關閉 zoom 才能用 FitToPages
            ps.FitToPagesWide = 1     # 寬度 = 1 頁
            ps.FitToPagesTall = False # 高度不限（可多頁）
            # 小邊距（inch）
            ps.LeftMargin = 0.3 * 72   # 0.3 inch（Excel 用 points，1 inch = 72 pt）
            ps.RightMargin = 0.3 * 72
            ps.TopMargin = 0.5 * 72
            ps.BottomMargin = 0.5 * 72
            ps.HeaderMargin = 0.2 * 72
            ps.FooterMargin = 0.2 * 72
            ps.CenterHorizontally = True
        except Exception:
            pass


def xlsx_list_to_pdf(xlsx_paths: list, progress_cb=None, batch_size: int = 50):
    """批次把 xlsx 轉成 PDF。為防 Excel 累積不穩定，每 batch_size 份重啟。

    輸出 PDF 會自動設為橫向 + 縮放到 1 頁寬，防止欄位被擠到下一頁。
    """
    try:
        import win32com.client
    except ImportError:
        return
    total = len(xlsx_paths)
    if total == 0:
        return

    def _new_excel():
        last_err = None
        for _ in range(3):
            try:
                e = win32com.client.DispatchEx("Excel.Application")
                try:
                    e.Visible = False
                except Exception:
                    pass
                try:
                    e.DisplayAlerts = False
                except Exception:
                    pass
                return e
            except Exception as err:
                last_err = err
                import time
                time.sleep(1)
        raise last_err if last_err else RuntimeError("無法啟動 Excel")

    excel = _new_excel()
    try:
        for idx, xlsx_path in enumerate(xlsx_paths):
            pdf_path = xlsx_path.rsplit(".", 1)[0] + ".pdf"
            try:
                wb = excel.Workbooks.Open(os.path.abspath(xlsx_path))
                _apply_pdf_pagesetup(wb)
                wb.ExportAsFixedFormat(0, os.path.abspath(pdf_path))
                wb.Close(SaveChanges=False)
            except Exception as e:
                if progress_cb:
                    try:
                        progress_cb(idx + 1, total, f"[WARN] {os.path.basename(xlsx_path)}: {e}")
                    except Exception:
                        pass
                continue

            if progress_cb:
                try:
                    progress_cb(idx + 1, total, os.path.basename(xlsx_path))
                except Exception:
                    pass

            if (idx + 1) % batch_size == 0 and (idx + 1) < total:
                try:
                    excel.Quit()
                except Exception:
                    pass
                excel = _new_excel()
    finally:
        try:
            excel.Quit()
        except Exception:
            pass


def read_raw_records(filepath: str) -> list:
    """讀取原始處方紀錄的所有資料列"""
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb["處方紀錄"]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            records.append(row)
    wb.close()
    return records
